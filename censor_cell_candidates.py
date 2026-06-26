"""Final cell-object censoring between raw Fiji outputs and PCA inputs."""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


SERVICE_COLUMNS = [
    "feature_row_id",
    "display_label",
    "image_name",
    "frame_id",
    "object_id",
    "object_type",
    "bbox_x",
    "bbox_y",
    "bbox_w",
    "bbox_h",
    "reference_roi_id",
    "inside_reference_roi",
    "reference_roi_overlap_fraction",
]

AUDIT_COLUMNS = [
    "feature_row_id",
    "display_label",
    "image_name",
    "frame_id",
    "object_id",
    "object_type",
    "bbox_x",
    "bbox_y",
    "bbox_w",
    "bbox_h",
    "reference_roi_id",
    "inside_reference_roi",
    "reference_roi_overlap_fraction",
    "object_pixels",
    "blue_pixels",
    "blue_pixel_percent",
    "censoring_status",
    "censoring_reason",
    "failed_stage",
    "last_passed_stage",
    "stage_trace",
    "selected_for_censored_frame",
]

FRAME_COLUMNS = [
    "feature_row_id",
    "display_label",
    "image_name",
    "frame_id",
    "object_id",
    "object_type",
    "selected_cell_count",
    "selected_object_pixels",
    "selected_blue_pixels",
    "selected_blue_pixel_fraction",
    "selected_blue_pixel_percent",
    "censoring_frame_status",
]

ALLOWED_THRESHOLD_OPS = {">", ">=", "<", "<=", "==", "!=", "between"}


@dataclass
class RowState:
    row: dict[str, str]
    status: str = "selected"
    reason: str = ""
    failed_stage: str = ""
    last_passed_stage: str = ""
    trace: list[str] = field(default_factory=list)

    @property
    def active(self) -> bool:
        return self.status == "selected"

    def reject(self, status: str, reason: str, stage_name: str, trace_text: str) -> None:
        self.status = status
        self.reason = reason
        self.failed_stage = stage_name
        self.trace.append(trace_text)

    def pass_stage(self, stage_name: str, trace_text: str) -> None:
        self.last_passed_stage = stage_name
        self.trace.append(trace_text)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Censor raw accepted cell candidates for downstream matrices/PCA.")
    parser.add_argument("--cell-features-raw", required=True, type=Path)
    parser.add_argument("--frame-features-raw", required=True, type=Path)
    parser.add_argument("--raw-legend", required=True, type=Path)
    parser.add_argument("--raw-overlay", required=True, type=Path)
    parser.add_argument("--config", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    return parser.parse_args()


def read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in fieldnames})


def load_config(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as handle:
        config = json.load(handle)
    if not isinstance(config, dict):
        raise ValueError("censoring config must be a JSON object")
    stages = config.get("selection_stages", [])
    if stages is None:
        stages = []
    if not isinstance(stages, list):
        raise ValueError("selection_stages must be a list")
    return config


def raw_accepted_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    if rows and "accepted_status" in rows[0]:
        return [
            row
            for row in rows
            if str(row.get("accepted_status", "")).strip().lower() == "accepted_cell_candidate"
        ]
    return rows


def stem_from_raw_path(path: Path, prefix: str, suffix: str) -> str:
    name = path.name
    if name.startswith(prefix) and name.endswith(suffix):
        return name[len(prefix) : -len(suffix)]
    return path.stem


def as_float(value: Any) -> float | None:
    try:
        if value is None or str(value).strip() == "":
            return None
        result = float(value)
        if math.isnan(result):
            return None
        return result
    except (TypeError, ValueError):
        return None


def as_int_like(value: Any) -> tuple[int, str]:
    text = str(value or "").strip()
    try:
        return int(float(text)), text
    except ValueError:
        return 999999999, text


def is_true(value: Any) -> bool:
    return str(value or "").strip().lower() in {"true", "1", "yes", "y"}


def validate_initial_rows(states: list[RowState], fieldnames: list[str]) -> None:
    missing = [column for column in SERVICE_COLUMNS if column not in fieldnames]
    for state in states:
        if missing:
            state.reject(
                "invalid_missing_required_column",
                "missing_required_columns:" + "|".join(missing),
                "precheck",
                "precheck:rejected_missing_required_columns",
            )
            continue
        if not is_true(state.row.get("inside_reference_roi")):
            state.reject(
                "invalid_not_inside_reference_roi",
                "inside_reference_roi_is_not_true",
                "precheck",
                "precheck:rejected_not_inside_reference_roi",
            )


def active_states(states: list[RowState]) -> list[RowState]:
    return [state for state in states if state.active]


def apply_keep_all(states: list[RowState], stage: dict[str, Any]) -> None:
    name = str(stage.get("name") or "keep_all")
    for state in active_states(states):
        state.reason = "keep_all"
        state.pass_stage(name, f"{name}:selected")


def apply_exclude_ids(states: list[RowState], stage: dict[str, Any]) -> None:
    name = str(stage.get("name") or "exclude_ids")
    id_column = str(stage.get("id_column") or "feature_row_id")
    ids = {str(value).strip() for value in stage.get("ids", [])}
    for state in active_states(states):
        value = str(state.row.get(id_column, "")).strip()
        if value in ids:
            state.reject("manual_excluded", f"{id_column}={value}", name, f"{name}:rejected_id_{value}")
        else:
            state.pass_stage(name, f"{name}:passed")


def apply_include_only_ids(states: list[RowState], stage: dict[str, Any]) -> None:
    name = str(stage.get("name") or "include_only_ids")
    id_column = str(stage.get("id_column") or "feature_row_id")
    ids = {str(value).strip() for value in stage.get("ids", [])}
    for state in active_states(states):
        value = str(state.row.get(id_column, "")).strip()
        if value in ids:
            state.pass_stage(name, f"{name}:passed")
        else:
            state.reject("include_only_filtered", f"{id_column}={value}", name, f"{name}:rejected_not_in_include_only")


def apply_top_n(states: list[RowState], stage: dict[str, Any]) -> None:
    name = str(stage.get("name") or "top_n")
    column = str(stage.get("column") or "")
    n = int(stage.get("n", 0))
    order = str(stage.get("order") or "desc").lower()
    current = active_states(states)
    if not column or n < 0:
        raise ValueError(f"{name}: top_n requires column and n >= 0")
    if len(current) <= n and str(stage.get("if_less_than_n", "")).lower() == "keep_all":
        for state in current:
            state.pass_stage(name, f"{name}:passed_keep_all_count_{len(current)}_le_{n}")
        return

    sortable: list[tuple[float, RowState]] = []
    for state in current:
        value = as_float(state.row.get(column))
        if value is None:
            state.reject(
                "invalid_missing_required_column",
                f"missing_or_invalid_numeric_column:{column}",
                name,
                f"{name}:rejected_missing_{column}",
            )
        else:
            sortable.append((value, state))

    reverse_primary = order != "asc"
    ranked = sorted(
        sortable,
        key=lambda item: (
            -item[0] if reverse_primary else item[0],
            as_int_like(item[1].row.get("object_id")),
            str(item[1].row.get("feature_row_id", "")),
        ),
    )
    keep_ids = {id(state) for _, state in ranked[:n]}
    ranks = {id(state): rank for rank, (_, state) in enumerate(ranked, start=1)}
    for _, state in ranked:
        rank = ranks[id(state)]
        if id(state) in keep_ids:
            state.pass_stage(name, f"{name}:passed_rank_{rank}_of_{n}")
        else:
            state.reject("not_in_top_n", f"{column}_rank_{rank}_over_{n}", name, f"{name}:rejected_rank_{rank}_over_{n}")


def threshold_passes(value: float, op: str, expected: Any) -> bool:
    if op == "between":
        if isinstance(expected, dict):
            low = as_float(expected.get("min"))
            high = as_float(expected.get("max"))
        elif isinstance(expected, list) and len(expected) == 2:
            low = as_float(expected[0])
            high = as_float(expected[1])
        else:
            low = high = None
        if low is None or high is None:
            raise ValueError("between threshold requires [min, max] or {min,max}")
        return low <= value <= high
    threshold = as_float(expected)
    if threshold is None:
        raise ValueError("threshold value must be numeric")
    if op == ">":
        return value > threshold
    if op == ">=":
        return value >= threshold
    if op == "<":
        return value < threshold
    if op == "<=":
        return value <= threshold
    if op == "==":
        return value == threshold
    if op == "!=":
        return value != threshold
    raise ValueError(f"unsupported threshold op: {op}")


def apply_threshold(states: list[RowState], stage: dict[str, Any]) -> None:
    name = str(stage.get("name") or "threshold")
    column = str(stage.get("column") or "")
    op = str(stage.get("op") or "")
    if not column or op not in ALLOWED_THRESHOLD_OPS:
        raise ValueError(f"{name}: threshold requires column and op in {sorted(ALLOWED_THRESHOLD_OPS)}")
    for state in active_states(states):
        value = as_float(state.row.get(column))
        if value is None:
            if str(stage.get("on_missing", "reject")).lower() == "keep":
                state.pass_stage(name, f"{name}:passed_missing_kept")
            else:
                state.reject(
                    "invalid_missing_required_column",
                    f"missing_or_invalid_numeric_column:{column}",
                    name,
                    f"{name}:rejected_missing_{column}",
                )
            continue
        if threshold_passes(value, op, stage.get("value")):
            state.pass_stage(name, f"{name}:passed")
        else:
            state.reject("threshold_failed", f"{column}_{op}_{stage.get('value')}", name, f"{name}:rejected_threshold_failed")


def apply_stages(states: list[RowState], config: dict[str, Any]) -> None:
    stages = config.get("selection_stages", [])
    if not stages and str(config.get("default_action", "")).lower() == "keep_all":
        stages = [{"name": "default_keep_all", "type": "keep_all"}]
    if not stages:
        raise ValueError("selection_stages is empty; set default_action=keep_all for no-op censoring")

    for stage in stages:
        stage_type = str(stage.get("type") or "").lower()
        if stage_type == "keep_all":
            apply_keep_all(states, stage)
        elif stage_type == "exclude_ids":
            apply_exclude_ids(states, stage)
        elif stage_type == "include_only_ids":
            apply_include_only_ids(states, stage)
        elif stage_type == "top_n":
            apply_top_n(states, stage)
        elif stage_type == "threshold":
            apply_threshold(states, stage)
        else:
            raise ValueError(f"unsupported selection stage type: {stage_type}")

    for state in active_states(states):
        if not state.reason:
            state.reason = "passed_selection_stages"


def audit_rows(states: list[RowState]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for state in states:
        row = {column: state.row.get(column, "") for column in AUDIT_COLUMNS}
        row.update(
            {
                "censoring_status": state.status,
                "censoring_reason": state.reason,
                "failed_stage": state.failed_stage,
                "last_passed_stage": state.last_passed_stage,
                "stage_trace": ";".join(state.trace),
                "selected_for_censored_frame": "true" if state.active else "false",
            }
        )
        rows.append(row)
    return rows


def selected_rows(states: list[RowState]) -> list[dict[str, str]]:
    return [state.row for state in states if state.active]


def numeric_sum(rows: list[dict[str, str]], column: str) -> float:
    return sum(as_float(row.get(column)) or 0.0 for row in rows)


def first_value(rows: list[dict[str, str]], column: str, fallback: str = "") -> str:
    for row in rows:
        value = str(row.get(column, "") or "").strip()
        if value:
            return value
    return fallback


def build_frame_row(selected: list[dict[str, str]], all_rows: list[dict[str, str]]) -> dict[str, Any]:
    frame_id = first_value(all_rows, "frame_id", "frame")
    image_name = first_value(all_rows, "image_name", "")
    object_pixels = numeric_sum(selected, "object_pixels")
    blue_pixels = numeric_sum(selected, "blue_pixels")
    fraction = blue_pixels / object_pixels if object_pixels else 0.0
    row: dict[str, Any] = {
        "feature_row_id": f"{frame_id}_censored_frame",
        "display_label": "",
        "image_name": image_name,
        "frame_id": frame_id,
        "object_id": "frame",
        "object_type": "censored_frame_cell_material",
        "selected_cell_count": len(selected),
        "selected_object_pixels": object_pixels,
        "selected_blue_pixels": blue_pixels,
        "selected_blue_pixel_fraction": fraction,
        "selected_blue_pixel_percent": 100.0 * fraction,
        "censoring_frame_status": "PASS" if selected else "FAIL_NO_CENSORED_OBJECTS",
    }
    add_weighted_columns(row, selected)
    return row


def add_weighted_columns(frame_row: dict[str, Any], rows: list[dict[str, str]]) -> None:
    total_weight = numeric_sum(rows, "object_pixels")
    if total_weight <= 0:
        return
    columns = set().union(*(row.keys() for row in rows)) if rows else set()
    for mean_column in sorted(column for column in columns if column.endswith("_mean")):
        prefix = mean_column[: -len("_mean")]
        std_column = f"{prefix}_std"
        weighted_values: list[tuple[float, float | None, float]] = []
        for row in rows:
            mean_value = as_float(row.get(mean_column))
            weight = as_float(row.get("object_pixels")) or 0.0
            if mean_value is not None and weight > 0:
                weighted_values.append((mean_value, as_float(row.get(std_column)), weight))
        if not weighted_values:
            continue
        weighted_mean = sum(mean * weight for mean, _, weight in weighted_values) / sum(weight for _, _, weight in weighted_values)
        frame_row[f"selected_weighted_{mean_column}"] = weighted_mean
        if std_column in columns:
            second_moment = sum((((std or 0.0) ** 2) + mean**2) * weight for mean, std, weight in weighted_values)
            second_moment /= sum(weight for _, _, weight in weighted_values)
            variance = max(0.0, second_moment - weighted_mean**2)
            frame_row[f"selected_weighted_{std_column}"] = math.sqrt(variance)


def write_blue_table(path: Path, selected: list[dict[str, str]], config: dict[str, Any]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    blue_config = config.get("blue_table", {}) if isinstance(config.get("blue_table", {}), dict) else {}
    sort_column = str(blue_config.get("sort_column") or "blue_pixel_percent")
    sort_order = str(blue_config.get("sort_order") or "desc").lower()
    add_sum_row = bool(blue_config.get("add_sum_row", True))
    rows = list(selected)
    rows.sort(
        key=lambda row: (
            -(as_float(row.get(sort_column)) or 0.0) if sort_order != "asc" else (as_float(row.get(sort_column)) or 0.0),
            as_int_like(row.get("object_id")),
            str(row.get("feature_row_id", "")),
        )
    )
    columns = [
        "image_name",
        "frame_id",
        "display_label",
        "feature_row_id",
        "object_type",
        "object_id",
        "object_pixels",
        "blue_pixels",
        "blue_pixel_fraction",
        "blue_pixel_percent",
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "blue_table_censored"
    worksheet.append(columns)
    for row in rows:
        worksheet.append([coerce_excel_value(row.get(column, "")) for column in columns])
    if add_sum_row:
        object_pixels = numeric_sum(rows, "object_pixels")
        blue_pixels = numeric_sum(rows, "blue_pixels")
        fraction = blue_pixels / object_pixels if object_pixels else 0.0
        worksheet.append([
            "SUM",
            first_value(rows, "frame_id"),
            "",
            "",
            "censored_selected_cell_material",
            "",
            object_pixels,
            blue_pixels,
            fraction,
            100.0 * fraction,
        ])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    if add_sum_row:
        for cell in worksheet[worksheet.max_row]:
            cell.font = Font(bold=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells) + 2
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(width, 10), 60)
    workbook.save(path)


def coerce_excel_value(value: Any) -> Any:
    numeric = as_float(value)
    if numeric is not None:
        return numeric
    return value


def draw_censored_overlay(raw_overlay: Path, output_path: Path, selected: list[dict[str, str]]) -> None:
    from PIL import Image, ImageDraw

    with Image.open(raw_overlay) as source:
        image = source.convert("RGB")
    draw = ImageDraw.Draw(image)
    for row in selected:
        x = int(as_float(row.get("bbox_x")) or 0)
        y = int(as_float(row.get("bbox_y")) or 0)
        w = max(1, int(as_float(row.get("bbox_w")) or 1))
        h = max(1, int(as_float(row.get("bbox_h")) or 1))
        label = str(row.get("display_label") or "")
        draw.rectangle((x, y, x + w, y + h), outline=(0, 255, 0), width=5)
        label_y = y + 4 if y < 24 else y - 22
        text_w = max(36, 8 * len(label) + 8)
        draw.rectangle((x, max(0, label_y - 2), x + text_w, max(18, label_y + 18)), fill=(0, 0, 0))
        draw.text((x + 4, max(0, label_y)), label, fill=(0, 255, 0))
    image.save(output_path, quality=90)


def output_fieldnames(input_fieldnames: list[str]) -> list[str]:
    fieldnames = list(input_fieldnames)
    for column in SERVICE_COLUMNS:
        if column not in fieldnames:
            fieldnames.append(column)
    return fieldnames


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    args = parse_args()
    for path in [args.cell_features_raw, args.frame_features_raw, args.raw_legend, args.raw_overlay, args.config]:
        if not path.exists():
            raise SystemExit(f"Input not found: {path}")

    config = load_config(args.config)
    cell_fieldnames, cell_rows = read_csv_rows(args.cell_features_raw)
    _, frame_rows = read_csv_rows(args.frame_features_raw)
    _, legend_rows = read_csv_rows(args.raw_legend)
    accepted = raw_accepted_rows(cell_rows)
    states = [RowState(dict(row)) for row in accepted]
    validate_initial_rows(states, cell_fieldnames)
    apply_stages(states, config)

    selected = selected_rows(states)
    image_stem = stem_from_raw_path(args.cell_features_raw, "cell_features_raw_", ".csv")
    args.output_dir.mkdir(parents=True, exist_ok=True)

    cell_output = args.output_dir / f"cell_features_censored_{image_stem}.csv"
    audit_output = args.output_dir / f"cell_censoring_audit_{image_stem}.csv"
    overlay_output = args.output_dir / f"cell_objects_overlay_censored_{image_stem}.jpg"
    frame_output = args.output_dir / f"frame_features_censored_{image_stem}.csv"
    blue_output = args.output_dir / f"blue_table_censored_{image_stem}.xlsx"

    write_csv(cell_output, output_fieldnames(cell_fieldnames), selected)
    write_csv(audit_output, AUDIT_COLUMNS, audit_rows(states))
    frame_row = build_frame_row(selected, accepted or legend_rows or frame_rows)
    frame_columns = list(FRAME_COLUMNS)
    for column in frame_row:
        if column not in frame_columns:
            frame_columns.append(column)
    write_csv(frame_output, frame_columns, [frame_row])
    write_blue_table(blue_output, selected, config)
    draw_censored_overlay(args.raw_overlay, overlay_output, selected)

    summary = {
        "cell_features_censored": str(cell_output),
        "cell_censoring_audit": str(audit_output),
        "cell_objects_overlay_censored": str(overlay_output),
        "frame_features_censored": str(frame_output),
        "blue_table_censored": str(blue_output),
        "raw_accepted_rows": len(accepted),
        "selected_rows": len(selected),
        "rejected_rows": len(accepted) - len(selected),
        "frame_status": frame_row["censoring_frame_status"],
    }
    (args.output_dir / "censoring_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
